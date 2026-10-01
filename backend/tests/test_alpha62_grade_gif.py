"""The 'Time to grade' slide always uses the bundled BIG Hat times_up.gif."""
import base64, hashlib, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import native_slides as ns

GIF = Path(__file__).resolve().parents[1] / "assets" / "slides" / "times_up.gif"


def _round(n_q, rtype="REG"):
    return ({"type": rtype, "name": "Short", "order": 1},
            {"round_type": rtype, "name": "Short",
             "questions": [{"number": i + 1, "question": f"Q{i+1}?", "answer": f"A{i+1}"} for i in range(n_q)]})


def _slides(n_q, rtype="REG"):
    ref, data = _round(n_q, rtype)
    return ns.render_round_section(data, ref)


def test_gif_is_bundled_and_real():
    assert GIF.is_file() and GIF.stat().st_size > 100_000
    assert GIF.read_bytes()[:6] in (b"GIF89a", b"GIF87a")


def test_grade_src_is_the_gif_not_the_svg():
    src = ns.grade_gif_src()
    assert src.startswith("data:image/gif;base64,")
    assert base64.b64decode(src.split(",", 1)[1]) == GIF.read_bytes()


def test_short_round_still_gets_grade_gif_after_review():
    for n in (3, 5, 10):
        slides = _slides(n)
        idx = [i for i, s in enumerate(slides) if (s.get("metadata") or {}).get("isGifStop")]
        assert len(idx) == 1, f"{n} questions: grade slide missing"
        el = slides[idx[0]]["elements"][0]
        assert el["src"].startswith("data:image/gif;base64,")
        rev = [i for i, s in enumerate(slides) if (s.get("metadata") or {}).get("isReview")]
        assert rev and rev[0] < idx[0], "grade slide must come after the review slide"


def test_every_round_type_has_it():
    for t in ("MC", "REG", "MISC", "MYS", "BIG"):
        assert any((s.get("metadata") or {}).get("isGifStop") for s in _slides(5, t)), t


def test_frozen_layout_lookup(tmp_path, monkeypatch):
    (tmp_path / "assets" / "slides").mkdir(parents=True)
    (tmp_path / "assets" / "slides" / "times_up.gif").write_bytes(b"GIF89a" + b"x" * 10)
    monkeypatch.setattr(sys, "_MEIPASS", str(tmp_path), raising=False)
    p = ns.bundled_asset_path("assets", "slides", "times_up.gif")
    assert p == tmp_path / "assets" / "slides" / "times_up.gif"


def test_build_script_locks_gif_in():
    s = (Path(__file__).resolve().parents[2] / "scripts" / "build_sidecar.py").read_text()
    assert "times_up.gif" in s and "LOCK-IN FAIL" in s
    assert "'assets'}{os.pathsep}assets" in s


def test_roundmaker_pptx_has_grade_gif_for_short_round():
    from routes import roundmaker as rm
    from pptx import Presentation
    import io
    for rtype in ("MC", "REG", "MISC", "MYS", "BIG"):
        qs = [{"number": i + 1, "question": f"Q{i+1}?", "answer": f"A{i+1}", "options": ["a", "b", "c", "d"]} for i in range(5)]
        prs = rm.generate_pptx({"round_type": rtype, "name": "Short", "questions": qs,
                                "tiebreaker": {"question": "t", "answer": "1"}}) if False else None
        prs = rm.generate_pptx({"round_type": rtype, "name": "Short", "questions": qs,
                                "tiebreaker": {"question": "t", "answer": "1"}})
        pres = prs if hasattr(prs, "slides") else Presentation(io.BytesIO(prs) if isinstance(prs, (bytes, bytearray)) else prs)
        pics = [sh for sl in pres.slides for sh in sl.shapes if sh.shape_type == 13]
        assert any(p.image.content_type == "image/gif" for p in pics), f"{rtype}: no grade gif in pptx"


# ---- LOCKED round slide ORDER (merchant-confirmed 2026-10-01) ----
def _kinds(rtype, n):
    ref, data = _round(n, rtype)
    data["tiebreaker"] = {"question": "TBQ", "answer": "TBA"}
    out = []
    for s in ns.render_round_section(data, ref):
        m = s.get("metadata") or {}
        if m.get("isTitleCard"): out.append("title")
        elif m.get("isGifStop"): out.append("gif")
        elif m.get("isReview"): out.append("review")
        elif m.get("isTiebreaker") and m.get("isAnswers"): out.append("tb-answer")
        elif m.get("isTiebreaker"): out.append("tb-question")
        elif m.get("isAnswers"): out.append("answers")
        else: out.append("question")
    return out


def test_mc_reg_misc_order_and_answers_last():
    for t in ("MC", "REG", "MISC"):
        k = _kinds(t, 5)
        assert k == ["title"] + ["question"] * 10 + ["review", "gif", "answers"], t


def test_mys_order_and_answers_last():
    assert _kinds("MYS", 5) == ["title"] + ["question"] * 9 + ["review", "gif", "answers"]


def test_big_order():
    assert _kinds("BIG", 1) == ["title", "question", "gif", "review", "answers", "tb-question", "tb-answer"]


def test_big_review_has_no_timer_in_player():
    src = (Path(__file__).resolve().parents[2] / "frontend/src/components/trivia/editor/PresentationMode.jsx").read_text()
    big = src[src.index("// BIG Round\n    if (roundType === 'BIG')"):]
    big = big[:big.index("return null;\n  }, [slides]);")]
    assert "relativeIndex === 1) return 300" in big          # only the question is timed
    assert "relativeIndex === 3" not in big                   # review (3) = no timer
